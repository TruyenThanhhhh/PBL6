import os
import uuid
import re
from pathlib import Path
from typing import List, Dict, Any

from app.config import CHUNK_SIZE, CHUNK_OVERLAP, UPLOAD_DIR
from app.database import get_db_connection

import json
import sqlite3
from app.services.embedding_service import embedding_service

class DocumentService:
    def process_and_store_document(self, doc_id: str, content: str, metadata: dict = None):
        # 1. Chia nhỏ văn bản (Chunking)
        chunks = self._chunk_text(content, chunk_size=512, overlap=50)
        
        # 2. Sinh Vector Embedding tự động qua BGE-M3
        embeddings = embedding_service.generate_embeddings(chunks)
        
        # 3. Lưu vào Database
        conn = sqlite3.connect("data/knowledge.db")
        cursor = conn.cursor()
        
        for idx, (chunk, emb) in enumerate(zip(chunks, embeddings)):
            cursor.execute(
                "INSERT INTO document_chunks (document_id, chunk_index, content, embedding, metadata) VALUES (?, ?, ?, ?, ?)",
                (doc_id, idx, chunk, json.dumps(emb), json.dumps(metadata or {}))
            )
            
        conn.commit()
        conn.close()
        return len(chunks)

    def _chunk_text(self, text: str, chunk_size: int = 512, overlap: int = 50):
        # Hàm phân tách văn bản đơn giản
        words = text.split()
        chunks = []
        for i in range(0, len(words), chunk_size - overlap):
            chunks.append(" ".join(words[i:i + chunk_size]))
        return chunks

document_service = DocumentService()

def extract_text_from_file(file_path: Path, file_type: str) -> List[Dict[str, Any]]:
    """
    Trích xuất text sạch từ file (PDF, DOCX, PPTX, TXT, MD) và giữ metadata số trang/slide.
    """
    results = []
    
    if file_type == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                if extracted.strip():
                    results.append({"page": i + 1, "text": extracted})
        except Exception as e:
            print(f"Lỗi đọc PDF: {e}")
            
    elif file_type == "docx":
        try:
            import docx
            doc = docx.Document(str(file_path))
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if full_text.strip():
                results.append({"page": 1, "text": full_text})
        except Exception as e:
            print(f"Lỗi đọc DOCX: {e}")

    elif file_type in ["pptx", "ppt"]:
        try:
            from pptx import Presentation
            prs = Presentation(str(file_path))
            for i, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    results.append({"page": i + 1, "text": "\n".join(slide_texts)})
        except Exception as e:
            print(f"Lỗi đọc PPTX: {e}")
            
    elif file_type in ["txt", "md", "py", "json", "csv"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                if content.strip():
                    results.append({"page": 1, "text": content})
        except Exception as e:
            print(f"Lỗi đọc text/md: {e}")
            
    return results

def recursive_chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """
    Chia nhỏ văn bản thông minh bám theo dấu câu và dòng trống
    """
    text = re.sub(r'\r\n', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    
    if len(text) <= chunk_size:
        return [text.strip()] if text.strip() else []
        
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunk = text[start:].strip()
            if chunk:
                chunks.append(chunk)
            break
            
        cut = -1
        for delim in ["\n\n", "\n", ". ", "; ", ", ", " "]:
            pos = text.rfind(delim, start + chunk_size // 2, end)
            if pos != -1:
                cut = pos + len(delim)
                break
                
        if cut == -1:
            cut = end
            
        chunk = text[start:cut].strip()
        if chunk:
            chunks.append(chunk)
        start = cut - overlap if cut - overlap > start else cut
        
    return chunks

def process_and_index_document(workspace_id: str, file_name: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Xử lý file tải lên: lưu trữ, trích xuất text, chunking và lưu vào SQLite + FTS
    """
    doc_id = str(uuid.uuid4())
    ext = file_name.split(".")[-1].lower() if "." in file_name else "txt"
    saved_path = UPLOAD_DIR / f"{doc_id}_{file_name}"
    
    with open(saved_path, "wb") as f:
        f.write(file_bytes)
        
    file_size = len(file_bytes)
    
    # 1. Trích xuất text sạch
    pages = extract_text_from_file(saved_path, ext)
    if not pages:
        # Nếu không trích xuất được text, cố gắng decode utf-8 an toàn (chỉ khi là file văn bản)
        if ext in ["txt", "md", "csv", "json"]:
            text = file_bytes.decode("utf-8", errors="ignore")
            if text.strip():
                pages = [{"page": 1, "text": text}]
                
    # 2. Chunking
    all_chunks = []
    chunk_index = 0
    for p in pages:
        page_chunks = recursive_chunk_text(p["text"])
        for c in page_chunks:
            chunk_id = str(uuid.uuid4())
            all_chunks.append({
                "id": chunk_id,
                "document_id": doc_id,
                "workspace_id": workspace_id,
                "chunk_index": chunk_index,
                "page_number": p["page"],
                "content": c
            })
            chunk_index += 1
            
    # 3. Lưu vào Database
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute("""
    INSERT INTO documents (id, workspace_id, file_name, file_type, file_size, file_path, total_chunks, status)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (doc_id, workspace_id, file_name, ext, file_size, str(saved_path), len(all_chunks), "indexed"))
    
    for c in all_chunks:
        cursor.execute("""
        INSERT INTO document_chunks (id, document_id, workspace_id, chunk_index, page_number, content)
        VALUES (?, ?, ?, ?, ?, ?)
        """, (c["id"], c["document_id"], c["workspace_id"], c["chunk_index"], c["page_number"], c["content"]))
        
        # Đánh chỉ mục Full-Text Search (FTS5)
        clean_fts_content = re.sub(r'["\']', '', c["content"])
        cursor.execute("""
        INSERT INTO chunks_fts (chunk_id, document_id, workspace_id, content)
        VALUES (?, ?, ?, ?)
        """, (c["id"], c["document_id"], c["workspace_id"], clean_fts_content))
        
    conn.commit()
    conn.close()
    
    return {
        "document_id": doc_id,
        "file_name": file_name,
        "file_type": ext,
        "total_chunks": len(all_chunks),
        "status": "indexed"
    }
